"""Build an editable 12-slide IND5003 presentation from committed result figures."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "PRESENTATION_DRAFT_2026-10-29.pptx"
FIG = ROOT / "outputs" / "fresh_course_experiment" / "final_eda_v3_shoe_fusion"

BG = RGBColor(14, 24, 33)
PANEL = RGBColor(27, 44, 57)
TEXT = RGBColor(240, 245, 247)
MUTED = RGBColor(177, 195, 201)
TEAL = RGBColor(53, 205, 178)
ORANGE = RGBColor(255, 177, 85)
RED = RGBColor(241, 112, 105)


def add_text(slide, text, x, y, w, h, *, size=20, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, bullets, x, y, w, h, *, size=17):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.08)
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)
    return box


def base(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)).fill.solid()
    top = slide.shapes[-1]
    top.fill.fore_color.rgb = TEAL
    top.line.fill.background()
    add_text(slide, title, 0.55, 0.35, 11.8, 0.5, size=27, bold=True)
    add_text(slide, f"IND5003 Vision II  |  {number:02d}", 10.2, 7.05, 2.5, 0.25, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_image(slide, filename, x, y, w, h):
    path = FIG / filename
    if path.is_file():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    else:
        add_text(slide, f"Missing local figure: {filename}", x, y, w, h, size=14, color=RED)


def add_metric(slide, label, value, x, y, color=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.2), Inches(1.0))
    shape.fill.solid(); shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = color
    add_text(slide, value, x + 0.1, y + 0.12, 2.0, 0.42, size=25, color=color, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.1, y + 0.58, 2.0, 0.24, size=10, color=MUTED, align=PP_ALIGN.CENTER)


def add_flow(slide):
    labels = ["Source\nvideo", "Event\nslice", "Pose +\nball", "Temporal\nidentity", "Geometry +\nphase", "Evidence\ntable"]
    for i, label in enumerate(labels):
        x = 0.55 + i * 2.08
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.7), Inches(1.55), Inches(1.1))
        sh.fill.solid(); sh.fill.fore_color.rgb = PANEL
        sh.line.color.rgb = TEAL if i in {0, 5} else ORANGE
        add_text(slide, label, x + 0.08, 2.95, 1.39, 0.55, size=17, bold=True, align=PP_ALIGN.CENTER)
        if i < len(labels) - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.58), Inches(3.02), Inches(0.42), Inches(0.4))
            ar.fill.solid(); ar.fill.fore_color.rgb = MUTED; ar.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = base(prs, "Penalty video analysis", 1)
    add_text(s, "Evidence-first Vision II pipeline", 0.7, 1.35, 8.0, 0.55, size=30, color=TEAL, bold=True)
    add_text(s, "Four 2018 World Cup shoot-outs · 39 penalty events", 0.72, 2.05, 7.0, 0.35, size=20, color=MUTED)
    add_metric(s, "contact candidates", "33/39", 8.3, 1.45)
    add_metric(s, "outcome fusion", "35/39", 10.75, 1.45, ORANGE)
    add_text(s, "Review draft · replace member names before recording", 0.72, 5.85, 7.0, 0.3, size=14, color=MUTED)

    s = base(prs, "Assignment outputs", 2)
    add_bullets(s, ["Nine requested event outputs", "Four camera-view intervals", "Feature value + source evidence + explicit unknown", "No numeric 90% threshold in the course brief"], 0.7, 1.35, 5.8, 4.8)
    add_text(s, "run-up / foot / steps / keeper movement / dive direction / dive height / ball trajectory / final ball position / goal-no-goal", 6.7, 1.7, 5.7, 1.7, size=22, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "behind kicker · 45° right · slow-motion kicker close-up · slow-motion behind goalkeeper", 6.8, 4.0, 5.5, 1.3, size=19, color=TEAL, align=PP_ALIGN.CENTER)

    s = base(prs, "Data and event definition", 3)
    add_metric(s, "penalty events", "39", 0.75, 1.45)
    add_metric(s, "matches", "4", 3.2, 1.45, ORANGE)
    add_metric(s, "keeper boxes", "835", 5.65, 1.45)
    add_bullets(s, ["Colombia–England 10 · Croatia–Denmark 10", "Spain–Russia 9 · Russia–Croatia 10", "Different views stay in one event partition", "24-row and core packages update existing events; they do not add independent penalties"], 0.85, 3.15, 11.6, 2.3)

    s = base(prs, "From video to evidence table", 4)
    add_flow(s)
    add_text(s, "Every output keeps a frame/time basis and provenance. Unsupported measurements abstain instead of becoming guesses.", 1.0, 5.15, 11.2, 0.7, size=22, color=MUTED, align=PP_ALIGN.CENTER)

    s = base(prs, "Labels, splits and validation", 5)
    add_bullets(s, ["Event-level partitions prevent views and frames from leaking across splits", "Human review covers foot, contact, movement, direction, height and outcome fields", "Eight-event frozen held-out check exists for the primary pipeline", "Newest ball/keeper detector arms were evaluated on exposed development groups"], 0.75, 1.35, 6.5, 4.9)
    add_text(s, "development ≠ independent test", 7.5, 2.2, 4.8, 0.6, size=29, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "Keep the denominator and evidence status visible.", 7.5, 3.15, 4.8, 0.7, size=22, color=TEXT, align=PP_ALIGN.CENTER)

    s = base(prs, "Exploratory coverage", 6)
    add_image(s, "01_fresh_coverage.png", 0.55, 1.25, 7.1, 5.25)
    add_metric(s, "pre-kick movement", "29/39", 8.2, 1.5)
    add_metric(s, "ball trajectory", "33/39", 10.7, 1.5, ORANGE)
    add_metric(s, "strict height", "0/39", 8.2, 3.0, RED)
    add_metric(s, "strict endpoint", "0/39", 10.7, 3.0, RED)
    add_text(s, "Coverage is not the same as accuracy.", 8.1, 4.8, 4.7, 0.55, size=21, color=MUTED, align=PP_ALIGN.CENTER)

    s = base(prs, "RQ1 · What can broadcast views show?", 7)
    add_image(s, "02_fresh_steps_timing.png", 0.55, 1.25, 7.0, 5.35)
    add_bullets(s, ["Temporal and pose evidence supports some contact and approach measurements", "Sparse ground projections are not a complete top-down trajectory", "Camera cuts and occlusion define the observable boundary"], 7.9, 1.55, 4.7, 4.2)

    s = base(prs, "RQ2 · Foot, steps and augmentation", 8)
    add_bullets(s, ["Foot evidence: about 13–14/39", "Step candidates: 32/39, but validation is development-only", "Ball augmentation: 65/91 → 62/91 core; 61/88 → 58/88 Arsenal", "Keeper photometric arm regressed on the external development group", "Neither new arm replaced the selected baseline"], 0.75, 1.35, 6.3, 4.9)
    add_text(s, "Model change must beat the frozen comparison, not just produce more boxes.", 7.6, 2.0, 4.7, 1.4, size=25, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)

    s = base(prs, "RQ3 · Ball path and outcome", 9)
    add_image(s, "03_fresh_ball_references.png", 0.55, 1.25, 5.9, 5.2)
    add_image(s, "05_fresh_scoreboard_outcomes.png", 6.7, 1.25, 5.9, 5.2)
    add_text(s, "33/39 ball-track candidates · 35/39 outcome-fusion candidates", 1.0, 6.6, 11.2, 0.35, size=18, color=MUTED, align=PP_ALIGN.CENTER)

    s = base(prs, "Failure boundaries", 10)
    add_bullets(s, ["Strict dive height: 0/39", "Strict event-final ball position: 0/39", "Provisional view: 17/39 image-plane height candidates", "Provisional view: 27/39 last-visible ball observations", "Provisional values are not accuracy and do not overwrite strict columns"], 0.8, 1.35, 6.2, 4.8)
    add_text(s, "unknown", 8.0, 1.8, 4.0, 0.8, size=42, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "is a valid result when the source evidence is insufficient", 7.55, 3.0, 4.9, 1.2, size=23, color=TEXT, align=PP_ALIGN.CENTER)

    s = base(prs, "Generalisation and reproducibility", 11)
    add_metric(s, "held-out primary runs", "8/8", 0.85, 1.45)
    add_metric(s, "held-out foot strict", "3/8", 3.35, 1.45, ORANGE)
    add_metric(s, "final focused tests", "53", 5.85, 1.45)
    add_bullets(s, ["Dell RTX 2060 portable package runs without human answers", "Runtime proves execution, not prediction accuracy", "Raw videos and weights stay outside the public repository", "Reports, manifests and hashes document what can be reproduced"], 0.85, 3.3, 11.2, 2.5)

    s = base(prs, "Conclusion and limitations", 12)
    add_bullets(s, ["The contribution is a reproducible, evidence-first pipeline with abstention", "Strongest coverage: contact, ball-track candidates, pre-kick movement and outcome fusion", "Main gaps: height, terminal ball position, full top-down path and camera playback", "Next data: endpoint boxes/timestamps, complete dive views and camera playback labels", "Finish the narrated video, report review and SharePoint package before submission"], 0.75, 1.25, 11.8, 4.7, size=20)
    add_text(s, "Thank you", 4.8, 6.1, 3.7, 0.5, size=30, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
