"""Score a completed deployment batch separately from all label-free inference."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import statistics


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_HUMAN = ROOT / "outputs/fresh_course_experiment/human_evaluation_2026-10-22/reference_source/penalty_annotations_39_human_v040.json"
DIRECTIONS = ("left", "center", "right")
HEIGHTS = ("low", "middle", "high")
SCOREBOARD_MAP = {"goal_likely": "goal", "no_goal_likely": "no_goal", "goal": "goal", "no_goal": "no_goal"}


def sha(path):
    with Path(path).open("rb") as stream:
        # Python 3.9 on the deployment host has no hashlib.file_digest.
        if hasattr(hashlib, "file_digest"):
            return hashlib.file_digest(stream, "sha256").hexdigest()
        digest = hashlib.sha256()
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
        return digest.hexdigest()


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def label(value):
    if isinstance(value, dict):
        if "labels" in value:
            values = value["labels"]
            return label(values[0]) if isinstance(values, list) and len(values) == 1 else None
        return label(value.get("label"))
    return "middle" if value == "mid" else value if isinstance(value, str) and value != "unknown" else None


def source_path(base, relative):
    relative = Path(relative)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError("Artifact path must stay inside its batch directory")
    target = (base / relative).resolve()
    if not target.is_relative_to(base.resolve()):
        raise ValueError("Artifact symlink escapes batch directory")
    return target


def load_completed_batch(directory):
    """Verify predictions and copied media before opening any reference file."""
    summary_path, manifest_path = directory / "summary.json", directory / "batch_manifest.json"
    summary, manifest = read(summary_path), read(manifest_path)
    if summary.get("complete") is not True or summary.get("failed_events") != 0:
        raise ValueError("Evaluate only after every requested prediction has completed successfully")
    digest = sha(manifest_path)
    if summary.get("batch_bindings_sha256") != digest:
        raise ValueError("Batch summary and manifest hashes differ")
    samples = manifest["bindings"]["samples"]
    expected = {s["core_id"]: s for s in samples}
    receipts = summary.get("events", [])
    if (len(expected) != len(samples) or len(receipts) != len(samples)
            or {r.get("core_id") for r in receipts} != set(expected)
            or summary.get("expected_events") != len(samples)):
        raise ValueError("Batch event identities or count differ from its manifest")
    predictions = {}
    bindings = []
    for receipt in receipts:
        cid = receipt["core_id"]
        sample = expected[cid]
        if (receipt.get("status") != "success" or receipt.get("batch_bindings_sha256") != digest
                or receipt.get("video_sha256") != sample["video_sha256"]):
            raise ValueError("Event receipt is unsuccessful or bound to a different batch/video")
        analysis_path = source_path(directory, receipt["analysis_path"])
        if sha(analysis_path) != receipt["analysis_sha256"]:
            raise ValueError("Prediction changed after its successful receipt")
        analysis = read(analysis_path)
        if (analysis.get("input_sha256") != sample["video_sha256"]
                or analysis.get("all_stages_executed_successfully") is not True
                or analysis.get("manual_annotations_used") is not False):
            raise ValueError("Prediction source, completion or annotation provenance is invalid")
        source = analysis["source"]
        filename = source["clip_filename"]
        if Path(filename).name != filename:
            raise ValueError("Invalid source video basename")
        media = source_path(analysis_path.parent, "source/" + filename)
        if source["clip_sha256"] != sample["video_sha256"] or sha(media) != sample["video_sha256"]:
            raise ValueError("Copied inference video differs from the bound input")
        predictions[cid] = analysis
        bindings.append({"core_id": cid, "video_sha256": sample["video_sha256"],
                         "analysis_sha256": receipt["analysis_sha256"], "analysis_path": receipt["analysis_path"]})
    return predictions, expected, {"batch_manifest_sha256": digest, "summary_sha256": sha(summary_path), "events": bindings}


def load_references(path, samples, human_path=None):
    from openpyxl import load_workbook
    document = read(path)
    if document.get("schema_version") != "core-contact-references-v1":
        raise ValueError("Expected the prepared core-contact-references-v1 labels.json")
    workbook_path = Path(document["source_workbook"])
    if sha(workbook_path) != document["source_workbook_sha256"]:
        raise ValueError("Human workbook changed after reference preparation")
    labels = {r["core_id"]: r for r in document["samples"]}
    if len(labels) != 10 or set(labels) != set(samples):
        raise ValueError("Evaluation requires exactly the same ten core events as the references")
    old = {e["clip_id"]: e for e in read(human_path)["events"]} if human_path else None
    book = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = book["\u8865\u680710\u6bb5"]
        original = book["\u539f\u6807\u6ce839\u6bb5"]
        hashes = {r[0]: r[16] for r in original.iter_rows(min_row=6, values_only=True) if r[0]}
        if sheet.cell(5, 7).value != "\u89e6\u7403\u65f6\u95f4(\u79d2)":
            raise ValueError("Workbook contact column schema changed")
        result = {}
        for row_number, row in enumerate(sheet.iter_rows(min_row=6, values_only=True), 6):
            if not row[0]:
                continue
            cid, clip_id = row[:2]
            if cid not in labels or cid in result:
                raise ValueError("Unexpected or duplicated workbook event")
            prepared, sample = labels[cid], samples[cid]
            if (prepared["clip_id"] != clip_id or sample.get("clip_id") != clip_id
                    or hashes.get(clip_id) != sample["video_sha256"]
                    or prepared["video_sha256"] != sample["video_sha256"]
                    or row_number != prepared["source_row"]
                    or row[6] != prepared["reference_contact_seconds"]):
                raise ValueError("Workbook, prepared reference and batch event bindings differ")
            inherited = {"foot": row[7], "ball_direction": row[10], "ball_height": label(row[11]),
                         "keeper_direction": row[14], "keeper_height": label(row[15])}
            if old is not None:
                old_event = old[clip_id]
                expected_old = {"foot": old_event["reviews"][0]["striking_foot"],
                    "ball_direction": old_event["ball_horizontal"], "ball_height": label(old_event["shot_height"]),
                    "keeper_direction": old_event["save_direction"], "keeper_height": label(old_event["save_height"])}
                if inherited != expected_old:
                    raise ValueError("Copied qualitative labels differ from the supplied human-39 source")
            outcome = {"\u8fdb\u7403": "goal", "\u6251\u51fa": "no_goal", "\u4e2d\u95e8\u6846": "no_goal", "\u504f\u51fa": "no_goal"}.get(row[17])
            if outcome is None or row[21] != "\u5df2\u5b8c\u6210\u590d\u6838":
                raise ValueError("Unrecognized or incomplete human outcome review")
            result[cid] = {"core_id": cid, "clip_id": clip_id, "video_sha256": sample["video_sha256"],
                "contact_seconds": row[6], "full_runup_steps": row[9], "keeper_movement_direction": row[12],
                "keeper_movement_seconds": row[13], "final_position_text": row[16], "outcome": outcome,
                "outcome_detail": row[17], "first_recorded_camera_cut_seconds": row[18],
                "direction_basis_raw": row[25], "direction_basis": "source_image_x_right" if
                    row[25] == "\u753b\u9762\u5de6\u53f3\uff08\u6cbf\u7528\u539f\u6807\u6ce8\uff09" else "unknown",
                "source_row": row_number, "contact_precision": prepared["reference_precision"],
                "inherited_qualitative_labels": True, **inherited,
                "ball_height_text_conflict": cid == "core_07",
                "ball_direction_text_conflict": cid == "core_08"}
    finally:
        book.close()
    if set(result) != set(samples):
        raise ValueError("Workbook does not cover the completed batch")
    return result, {"prepared_reference_sha256": sha(path), "workbook_sha256": sha(workbook_path),
        "human39_reference_sha256": sha(human_path) if human_path else None,
        "references_read_only_after_prediction_integrity_checks": True}


def categorical_metrics(rows, classes):
    eligible = [r for r in rows if r.get("comparable", True) and r.get("reference") in classes]
    available = [r for r in eligible if r.get("prediction") in classes]
    correct = sum(r["prediction"] == r["reference"] for r in available)
    per_class = {}
    for category in classes:
        tp = sum(r.get("prediction") == r["reference"] == category for r in eligible)
        fp = sum(r.get("prediction") == category and r["reference"] != category for r in eligible)
        fn = sum(r["reference"] == category and r.get("prediction") != category for r in eligible)
        per_class[category] = {"references": tp + fn, "true_positive": tp, "false_positive": fp,
            "false_negative_including_unknown": fn, "recall": tp / (tp + fn) if tp + fn else None,
            "f1": 2 * tp / (2 * tp + fp + fn) if 2 * tp + fp + fn else None}
    scores = [p["f1"] for p in per_class.values() if p["f1"] is not None]
    return {"batch_events": len(rows), "comparable_reference_events": len(eligible),
        "prediction_available_events": len(available), "correct_events": correct,
        "coverage_on_comparable_references": len(available) / len(eligible) if eligible else None,
        "strict_accuracy_unknown_as_incorrect": correct / len(eligible) if eligible else None,
        "conditional_accuracy_on_available": correct / len(available) if available else None,
        "macro_f1_on_represented_classes_unknown_as_false_negative": statistics.mean(scores) if scores else None,
        "per_class": per_class, "excluded_comparisons": [{"core_id": r["core_id"], "reason": r.get("reason", "unknown_reference")}
                                                       for r in rows if r not in eligible], "rows": rows}


def numeric_metrics(rows, tolerances):
    eligible = [r for r in rows if r.get("comparable", True) and number(r.get("reference"))]
    available = [r for r in eligible if number(r.get("prediction"))]
    signed = [r["prediction"] - r["reference"] for r in available]
    errors = [abs(x) for x in signed]
    return {"batch_events": len(rows), "comparable_reference_events": len(eligible),
        "prediction_available_events": len(available),
        "coverage_on_comparable_references": len(available) / len(eligible) if eligible else None,
        "mae_on_available": statistics.mean(errors) if errors else None,
        "median_absolute_error_on_available": statistics.median(errors) if errors else None,
        "maximum_absolute_error_on_available": max(errors) if errors else None,
        "mean_signed_prediction_minus_reference": statistics.mean(signed) if signed else None,
        "tolerances": {str(t): {"correct_events": sum(e <= t + 1e-9 for e in errors),
            "strict_hit_fraction_unknown_as_miss": sum(e <= t + 1e-9 for e in errors) / len(eligible) if eligible else None,
            "conditional_hit_fraction": sum(e <= t + 1e-9 for e in errors) / len(available) if available else None} for t in tolerances},
        "excluded_comparisons": [{"core_id": r["core_id"], "reason": r.get("reason", "unknown_reference")}
                                 for r in rows if r not in eligible], "rows": rows}


def evaluate(predictions, references):
    categories = {name: [] for name in ("kicking_foot", "keeper_dive_direction", "goal_outcome",
        "keeper_dive_height", "keeper_ordinal_height_auxiliary", "ball_height_auxiliary", "ball_direction_auxiliary")}
    numbers = {name: [] for name in ("contact_seconds", "full_runup_steps", "first_camera_cut_diagnostic")}
    event_rows = []
    for cid, analysis in sorted(predictions.items()):
        ref = references[cid]
        fields, auxiliary = analysis["required_outputs"], analysis.get("auxiliary_models", {})
        contact = analysis.get("trained_contact") or {}
        foot = label(fields.get("kicking_foot"))
        categories["kicking_foot"].append({"core_id": cid, "reference": ref["foot"], "prediction": foot})
        direction = fields.get("keeper_dive_direction")
        image_direction = {"image_left": "left", "image_right": "right", "center": "center", "unknown": None}.get(direction)
        direction_compatible = ref["direction_basis"] == "source_image_x_right" and direction in ("image_left", "image_right", "center", "unknown", None)
        categories["keeper_dive_direction"].append({"core_id": cid, "reference": ref["keeper_direction"],
            "prediction": image_direction, "comparable": direction_compatible,
            "reason": None if direction_compatible else "direction_coordinate_conventions_not_matched"})
        categories["goal_outcome"].append({"core_id": cid, "reference": ref["outcome"],
            "prediction": SCOREBOARD_MAP.get(fields.get("goal_outcome"))})
        numbers["contact_seconds"].append({"core_id": cid, "reference": ref["contact_seconds"],
            "prediction": (contact.get("prediction") or {}).get("time_seconds"),
            "reference_precision": ref["contact_precision"]})
        shoes = auxiliary.get("trained_shoes") or {}
        steps = fields.get("visible_runup_steps")
        scope_matches = shoes.get("count_complete") is True and shoes.get("full_approach_visible") is True
        numbers["full_runup_steps"].append({"core_id": cid, "reference": ref["full_runup_steps"],
            "prediction": steps, "comparable": scope_matches,
            "reason": None if scope_matches else "human_full_approach_vs_automatic_observed_window",
            "unaligned_difference_diagnostic_only": steps - ref["full_runup_steps"] if number(steps) else None})
        window = contact.get("automatic_window") or {}
        fps = analysis["source"].get("fps")
        first_end = window.get("end_frame_exclusive") / fps if number(window.get("end_frame_exclusive")) and number(fps) and fps > 0 else None
        numbers["first_camera_cut_diagnostic"].append({"core_id": cid,
            "reference": ref["first_recorded_camera_cut_seconds"], "prediction": first_end,
            "scope": "first_automatic_shot_end_not_four_camera_class_accuracy"})
        keeper_height = ((auxiliary.get("trained_keeper") or {}).get("prediction") or {}).get("height_model_candidate")
        presented_height = fields.get("keeper_dive_height")
        qualitative = isinstance(presented_height, dict) and presented_height.get("definition") == "human_qualitative_image_plane"
        categories["keeper_dive_height"].append({"core_id": cid, "reference": ref["keeper_height"],
            "prediction": label(presented_height.get("value")) if qualitative else None,
            "comparable": qualitative,
            "reason": None if qualitative else "legacy_strict_height_is_a_different_target"})
        categories["keeper_ordinal_height_auxiliary"].append({"core_id": cid,
            "reference": ref["keeper_height"], "prediction": label(keeper_height),
            "scope": "inherited_human_ordinal_category_not_physical_or_strict_phase_height"})
        ball = (auxiliary.get("trained_ball") or {}).get("predictions") or {}
        categories["ball_height_auxiliary"].append({"core_id": cid, "reference": ref["ball_height"],
            "prediction": label(ball.get("ball_height_category")), "comparable": not ref["ball_height_text_conflict"],
            "reason": "core07_old_mid_vs_new_low_text_conflict" if ref["ball_height_text_conflict"] else None})
        categories["ball_direction_auxiliary"].append({"core_id": cid, "reference": ref["ball_direction"],
            "prediction": label(ball.get("ball_direction")), "comparable": False,
            "reason": "external_head_behind_kicker_or_undeclared_convention_vs_core_image_coordinates"})
        top_down = fields.get("runup_top_down_trajectory") or {}
        event_rows.append({"core_id": cid, "clip_id": ref["clip_id"], "video_sha256": ref["video_sha256"],
            "end_to_end_seconds": analysis.get("end_to_end_seconds", analysis.get("summed_stage_runtime_seconds")),
            "video_duration_seconds": analysis["source"].get("slice_end_sec"),
            "selected_output_sources": analysis.get("selected_output_sources", {}),
            "top_down_candidate_points": len(top_down.get("points", [])),
            "complete_top_down_trajectory_claimed": top_down.get("complete_trajectory_verified") is True,
            "ball_trajectory_points": len(fields.get("ball_trajectory") or []),
            "ball_final_position_available": final_position_value(fields.get("ball_final_position")) is not None,
            "keeper_pre_movement_output": fields.get("keeper_pre_kick_movement"),
            "keeper_pre_movement_reference": ref["keeper_movement_direction"],
            "keeper_pre_movement_comparable": False,
            "keeper_pre_movement_reason": "predicted_movement_level_vs_human_direction_or_start_time",
            "strict_keeper_height_output": strict_height_value(presented_height),
            "strict_keeper_height_comparable": False,
            "strict_keeper_height_reason": "ordinal_reference_does_not_define_strict_torso_clearance_metric",
            "camera_segment_count": len(analysis.get("camera_intervals") or []),
            "manual_four_camera_class_interval_ground_truth_available": False,
            "final_position_reference_text": ref["final_position_text"],
            "final_position_coordinate_ground_truth_available": False})
    categorical = {name: categorical_metrics(rows, ("left", "right") if name == "kicking_foot" else
        ("goal", "no_goal") if name == "goal_outcome" else HEIGHTS if "height" in name else DIRECTIONS)
        for name, rows in categories.items()}
    numeric = {name: numeric_metrics(rows, (0, 1) if "steps" in name else (0.1, 0.2, 0.5))
               for name, rows in numbers.items()}
    runtime = [row["end_to_end_seconds"] for row in event_rows if number(row["end_to_end_seconds"])]
    ratios = [r["end_to_end_seconds"] / r["video_duration_seconds"] for r in event_rows
              if number(r["end_to_end_seconds"]) and number(r["video_duration_seconds"]) and r["video_duration_seconds"] > 0]
    checks = {
        "kicker_top_down_candidate_points": lambda a: bool((a["required_outputs"].get("runup_top_down_trajectory") or {}).get("points")),
        "kicking_foot": lambda a: a["required_outputs"].get("kicking_foot") in ("left", "right"),
        "observed_window_step_count": lambda a: number(a["required_outputs"].get("visible_runup_steps")),
        "keeper_pre_kick_movement_level": lambda a: label(a["required_outputs"].get("keeper_pre_kick_movement")) is not None,
        "keeper_dive_direction": lambda a: a["required_outputs"].get("keeper_dive_direction") in ("image_left", "center", "image_right"),
        "keeper_strict_height_candidate": lambda a: label(strict_height_value(a["required_outputs"].get("keeper_dive_height"))) in HEIGHTS,
        "keeper_qualitative_height_candidate": lambda a: isinstance(a["required_outputs"].get("keeper_dive_height"), dict)
            and label(a["required_outputs"]["keeper_dive_height"].get("value")) in HEIGHTS,
        "ball_trajectory_candidate": lambda a: bool(a["required_outputs"].get("ball_trajectory")),
        "ball_final_position": lambda a: final_position_value(a["required_outputs"].get("ball_final_position")) is not None,
        "goal_outcome": lambda a: a["required_outputs"].get("goal_outcome") in SCOREBOARD_MAP,
        "camera_interval_candidates": lambda a: bool(a.get("camera_intervals")),
    }
    coverage = {name: {"output_events": sum(bool(check(a)) for a in predictions.values()),
                      "all_events": len(predictions),
                      "fraction": sum(bool(check(a)) for a in predictions.values()) / len(predictions),
                      "availability_is_not_accuracy": True} for name, check in checks.items()}
    return {"output_coverage": coverage, "categorical": categorical, "numeric": numeric, "events": event_rows,
        "runtime": {"events": len(runtime), "median_seconds": statistics.median(runtime) if runtime else None,
            "p95_seconds": statistics.quantiles(runtime, n=20, method="inclusive")[18] if len(runtime) > 1 else runtime[0] if runtime else None,
            "maximum_seconds": max(runtime) if runtime else None,
            "median_real_time_ratio": statistics.median(ratios) if ratios else None},
        "unscored_teacher_outputs": {
            "top_down_runup_trajectory": "no independent corresponding path or calibration ground truth",
            "keeper_pre_kick_movement": "output level and reference direction/start time are different targets",
            "keeper_dive_height_strict": "strict metric differs from inherited human ordinal category",
            "ball_trajectory": "no independent ball-centre sequence",
            "ball_final_position": "free-text result/location has no terminal coordinate/time ground truth",
            "four_camera_intervals": "one recorded cut per core clip is not four-class interval annotation"}}


def final_position_value(value):
    return value.get("event_final_position") if isinstance(value, dict) else value


def strict_height_value(value):
    return value.get("strict_mask_pose_category") if isinstance(value, dict) else value


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch-dir", type=Path, required=True)
    parser.add_argument("--references", type=Path, required=True, help="Prepared core data/labels.json")
    parser.add_argument("--human-reference", type=Path, default=DEFAULT_HUMAN)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise FileExistsError("Choose a new evaluation directory; earlier reports are preserved")
    predictions, samples, batch_bindings = load_completed_batch(args.batch_dir.resolve())
    references, reference_bindings = load_references(args.references, samples, args.human_reference)
    report = {"schema_version": "offline-video-deployment-evaluation-v1",
        "evaluation_kind": "deployment_smoke_on_previously_trained_development_events",
        "held_out_test": False, "out_of_fold_predictions": False, "may_contain_in_sample_predictions": True,
        "all_predictions_finished_before_reference_read": True,
        "batch_bindings": batch_bindings, "reference_bindings": reference_bindings,
        "code_sha256": sha(Path(__file__)), **evaluate(predictions, references),
        "limitations": ["Ten events overlap training/development; these are not generalization estimates.",
            "Unknown predictions count as misses for strict metrics; conditional scores disclose their smaller denominator.",
            "Incompatible target definitions and coordinate systems are excluded, not counted as correct or silently transformed.",
            "Contact times are coarse human observations, not exact physical contact frames.",
            "Inherited single-human qualitative categories are not independent new reviews.",
            "Output availability and code completion do not establish trajectory or endpoint accuracy."]}
    args.output.mkdir(parents=True)
    (args.output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"report": str(args.output / "report.json"), "events": len(predictions),
        "held_out_test": False, "categorical": {k: {a: b for a, b in v.items() if a not in ("rows", "per_class", "excluded_comparisons")}
                                               for k, v in report["categorical"].items()}}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
