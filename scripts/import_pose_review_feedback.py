"""Import the 2026-10-27 pose review as source-bound observations, not labels."""

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import cv2
import openpyxl


ROOT = Path(__file__).resolve().parents[1]
MODEL = "YOLO11m-pose"
KEEPER_SHEET = "人工反馈"
SHOOTER_SHEET = "射手关节异常"
HEADERS = {
    KEEPER_SHEET: (
        "视频", "检查对象", "定位帧范围", "人工观察", "问题归类", "确认范围与歧义", "后续处理",
    ),
    SHOOTER_SHEET: (
        "视频", "射手ID", "起始帧", "结束帧", "恢复帧", "身体部位", "用户目视观察",
        "可见性与原因", "记录结论及范围", "补充复核备注（可填）",
    ),
}
KEEPER_IDS = ("05-04", "07-08", "12-02", "02-06", "02-05", "10-01", "17-08", "19-08")
# These dated, free-text rows require an explicit contract before normalization.
KEEPER_CONTRACT = {
    "05-04": ("门将，ID0", "人工：67帧后；模型索引：67–79", "67帧后ID0丢失。"),
    "07-08": (
        "门将ID2；射手ID1右手遮挡视角", "67；68帧后",
        "用户更正：第67帧门将ID2因遮挡无法判断；68帧后正常。",
    ),
    "12-02": ("前景新人物与原射手的关联检查", "此前重点145–170", "一切正常。"),
    "02-06": (
        "门将，ID1/2/3", "107–108；124；128；其他帧未稳定检测",
        "新出现的全是同一个人，都是门将；没有长时间稳定跟踪。",
    ),
    **{clip: (
        "门将，未分配稳定模型ID", f"全段0–{end}",
        "门将肉眼明确、动作可见，但全程无YOLO框。",
    ) for clip, end in (("02-05", 140), ("10-01", 155), ("17-08", 125), ("19-08", 128))},
}
SHOOTER_ROWS = (
    ("05-04", 7, 51, 56, 57), ("05-04", 7, 60, 60, None),
    ("05-04", 7, 65, 65, None), ("10-01", 1, 74, 74, None),
    ("10-01", 1, 79, 79, None), ("10-01", 1, 80, 83, None),
    ("10-01", 1, 113, 114, None), ("17-08", 1, 77, 77, None),
    ("17-08", 1, 79, 79, None), ("19-08", 0, 29, 29, None),
    ("19-08", 0, 38, 40, None), ("19-08", 0, 57, 66, None),
)
FULL_MISSES = {"02-05": 140, "10-01": 155, "17-08": 125, "19-08": 128}


def sha256(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def binding():
    return {
        "reviewed_model": MODEL,
        "status": "binding_pending",
        "reviewed_inference_output_sha256": None,
        "id_scope": "video_and_exact_reviewed_model_output_only",
        "legacy_output_id_mapping_allowed": False,
    }


def usage_constraints():
    return {
        "evidence_level": "human_reported_visual_observation",
        "training_label_eligible": False,
        "human_coordinate_ground_truth": False,
        "whole_clip_pose_accuracy_verified": False,
        "independent_consensus": False,
        "allowed_uses": ["source_video_frame_review", "pending_model_regression_review"],
        "forbidden_uses": [
            "coordinate_training_labels", "striking_foot_ground_truth", "ball_contact_ground_truth",
            "unbound_track_id_join", "held_out_accuracy_claim",
        ],
    }


def video_metadata(path):
    path = Path(path).resolve(strict=True)
    digest = sha256(path)
    capture = cv2.VideoCapture(str(path))
    try:
        count = capture.get(cv2.CAP_PROP_FRAME_COUNT)
        fps = capture.get(cv2.CAP_PROP_FPS)
        if not capture.isOpened() or count <= 0 or not count.is_integer() or fps <= 0:
            raise ValueError(f"Unreadable video metadata: {path}")
    finally:
        capture.release()
    return {
        "video_id": path.stem, "path": str(path), "sha256": digest,
        "size_bytes": path.stat().st_size, "frame_count": int(count), "fps": fps,
        "metadata_method": "opencv_container_metadata",
    }


def read_source(path):
    """Read cells without saving, evaluating formulas, or rewriting the workbook."""
    path = Path(path).resolve(strict=True)
    digest = sha256(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != [KEEPER_SHEET, SHOOTER_SHEET]:
            raise ValueError("Unexpected source sheets; this importer accepts the dated review schema only")
        snapshots = {}
        for sheet in workbook:
            rows = {}
            for number, cells in enumerate(sheet.iter_rows(), start=1):
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"Formulas are not accepted as source observations: {sheet.title}:{number}")
                values = [cell.value for cell in cells]
                if any(value is not None for value in values):
                    if any(value is not None and not isinstance(value, (str, int, float, bool)) for value in values):
                        raise ValueError(f"Unexpected source cell type: {sheet.title}:{number}")
                    rows[number] = values
            width = len(HEADERS[sheet.title])
            if tuple(rows.get(6, [])[:width]) != HEADERS[sheet.title]:
                raise ValueError(f"Unexpected source headers: {sheet.title}")
            snapshots[sheet.title] = rows
        if MODEL not in str(snapshots[KEEPER_SHEET].get(2, [None])[0]):
            raise ValueError("The feedback must explicitly refer to YOLO11m-pose")
        for title in HEADERS:
            if not any("从0开始" in str(value) for row in snapshots[title].values() for value in row):
                raise ValueError(f"Missing zero-based source frame convention: {title}")
        keeper = [snapshots[KEEPER_SHEET].get(number, [])[:7] for number in range(7, 15)]
        shooter = [snapshots[SHOOTER_SHEET].get(number, [])[:10] for number in range(7, 19)]
        if tuple(row[0] if row else None for row in keeper) != KEEPER_IDS:
            raise ValueError("Expected eight ordered keeper/person observations")
        if tuple(tuple(row[:5]) for row in shooter) != SHOOTER_ROWS:
            raise ValueError("Expected the twelve supplied shooter ranges and model IDs")
        for row in keeper:
            if len(row) != 7 or any(not isinstance(value, str) or not value.strip() for value in row):
                raise ValueError("Keeper/person source observations must be complete text rows")
            if tuple(row[1:4]) != KEEPER_CONTRACT[row[0]]:
                raise ValueError(f"Keeper source wording changed; normalization requires review: {row[0]}")
        for row in shooter:
            if len(row) != 10 or any(not isinstance(value, str) or not value.strip() for value in row[5:9]):
                raise ValueError("Shooter source observations must preserve all descriptive fields")
            if any(type(value) is not int for value in row[1:4]) or (row[4] is not None and type(row[4]) is not int):
                raise ValueError("Model IDs and frame indices must be integers")
        for title, end in ((KEEPER_SHEET, 14), (SHOOTER_SHEET, 18)):
            if any(number > end and (title == SHOOTER_SHEET or row[0] in KEEPER_IDS)
                   for number, row in snapshots[title].items()):
                raise ValueError(f"Unexpected additional observation row in {title}")
    finally:
        workbook.close()
    if sha256(path) != digest:
        raise ValueError("Source workbook changed during import")
    return {
        "path": str(path), "sha256": digest, "size_bytes": path.stat().st_size,
        "opened_read_only": True, "source_unchanged": True,
        "received_date_as_declared": "2026-10-27", "actual_review_date": None,
        "reviewer_identity_independently_verified": False,
        "sheets": [
            {"name": title, "nonempty_rows": [{"row": number, "cells": values} for number, values in rows.items()]}
            for title, rows in snapshots.items()
        ],
    }, keeper, shooter


def source_reference(source, sheet, number, values):
    return {
        "workbook_sha256": source["sha256"], "sheet": sheet, "row": number,
        "raw_cells": dict(zip(HEADERS[sheet], values)),
    }


def frame_range(start, end, video):
    if type(start) is not int or type(end) is not int or not 0 <= start <= end < video["frame_count"]:
        raise ValueError(f"Invalid zero-based inclusive frame range for {video['video_id']}: {start}-{end}")
    return {"start": start, "end": end, "inclusive": True, "index_base": 0}


def base_record(source, sheet, number, values, video):
    return {
        "observation_id": f"{'keeper' if sheet == KEEPER_SHEET else 'shooter'}-row-{number}",
        "source": source_reference(source, sheet, number, values),
        "source_video": video, "model_binding": binding(),
        "usage_constraints": usage_constraints(),
        "coordinate_labels": None, "striking_foot_label": None, "ball_contact_frame_label": None,
        "review_of_current_model_status": "not_evaluated",
    }


def keeper_observations(source, rows, videos):
    results = []
    for number, values in enumerate(rows, start=7):
        clip = values[0]
        video = videos[clip]
        record = base_record(source, KEEPER_SHEET, number, values, video)
        record.update({
            "reported_observation": values[3], "reported_category": values[4],
            "reported_scope_and_ambiguity": values[5], "reported_next_action": values[6],
            "identity_switch_confirmed": False, "whole_clip_pass": False,
            "whole_clip_review_scope": clip in FULL_MISSES,
        })
        if clip == "05-04":
            ranges = [(67, 79)]
            record.update({
                "issue_kind": "late_track_loss", "subject_model_ids": [0],
                "frame_range_basis": "Source model-index range; human wording says after frame 67",
                "identity_relationship": "not_determined", "reviewed_state": "reported_missing",
            })
        elif clip == "07-08":
            ranges = [(67, 67), (68, video["frame_count"] - 1)]
            record.update({
                "issue_kind": "occlusion_unknown_then_reported_normal", "subject_model_ids": [2],
                "context_model_ids": [1], "superseded_incorrect_keeper_id": 0,
                "frame_range_states": ["unknown", "reported_normal"],
                "frame_range_basis": "Frame 67 and from 68 to source-video end; corrected keeper ID is 2",
                "identity_relationship": "not_determined", "reviewed_state": "partly_unknown",
            })
        elif clip == "12-02":
            ranges = [(145, 170)]
            record.update({
                "issue_kind": "local_person_association_reported_normal", "subject_model_ids": [],
                "frame_range_basis": "Previously highlighted association check only",
                "identity_relationship": "local_association_reported_normal",
                "reviewed_state": "reported_normal", "goalkeeper_coverage_confirmed": False,
            })
        elif clip == "02-06":
            ranges = [(107, 108), (124, 124), (128, 128)]
            record.update({
                "issue_kind": "same_person_track_fragmentation", "subject_model_ids": [1, 2, 3],
                "frame_range_basis": "Reported visible fragments; other frames have unstable detection",
                "individual_id_to_range_assignment": None,
                "identity_relationship": "same_goalkeeper_reported", "reviewed_state": "fragmented",
            })
        else:
            end = FULL_MISSES[clip]
            if end != video["frame_count"] - 1:
                raise ValueError(f"Whole-clip source range differs from video length: {clip}")
            ranges = [(0, end)]
            record.update({
                "issue_kind": "whole_clip_goalkeeper_detection_missing", "subject_model_ids": [],
                "frame_range_basis": "Whole-clip range explicitly stated in source",
                "identity_relationship": "no_stable_model_id", "reviewed_state": "reported_missing",
                "human_reports_visible_goalkeeper_motion": True,
            })
        record["review_frame_ranges"] = [frame_range(start, end, video) for start, end in ranges]
        results.append(record)
    return results


def shooter_observations(source, rows, videos):
    results = []
    for number, values in enumerate(rows, start=7):
        clip, model_id, start, end, recovery, body, observation, visibility, conclusion, note = values
        video = videos[clip]
        record = base_record(source, SHOOTER_SHEET, number, values, video)
        if recovery is not None:
            frame_range(recovery, recovery, video)
            if recovery <= end:
                raise ValueError("Recovery must follow the reported issue interval")
        record.update({
            "subject_model_id": model_id, "subject_role_as_confirmed": "shooter",
            "issue_frame_range": frame_range(start, end, video), "reported_recovery_frame": recovery,
            "body_region_as_reported": body, "reported_observation": observation,
            "reported_visibility_and_cause": visibility, "reported_scope": conclusion,
            "additional_review_note": note,
            "range_basis": "start_and_recovery_implied_interval" if recovery is not None else "explicit_source_range",
            "visible_skeleton_missing_reported": "丢失" in observation,
            "left_right_channel_status": "ambiguous_unknown" if "混淆" in observation else "not_assessed",
            "actual_left_right_swap_confirmed": False,
            "csv_coordinate_missing_confirmed": False,
            "specific_missing_joint_indices": None,
            "corrected_coordinates_available": False,
        })
        results.append(record)
    return results


def regression_checklist(records):
    return {
        "schema_version": "pose-review-regression-checklist-v1",
        "status": "pending", "model_binding": binding(),
        "model_regression_passed": None, "accuracy_computed": False,
        "reviewed_items": 0,
        "required_before_model_comparison": [
            "Freeze and hash the exact reviewed YOLO11m-pose output",
            "Verify the output source-video hashes and zero-based frame convention",
            "Bind model IDs within that output; never join legacy IDs by number alone",
            "Review candidate output in the listed source video and frame ranges",
        ],
        "items": [{
            "observation_id": record["observation_id"],
            "source": record["source"],
            "source_video_sha256": record["source_video"]["sha256"],
            "video_id": record["source_video"]["video_id"],
            "frame_ranges": record.get("review_frame_ranges", [record.get("issue_frame_range")]),
            "status": "pending", "result": None, "candidate_output_sha256": None,
            "training_label_eligible": False,
        } for record in records],
    }


def write_json(path, payload):
    with Path(path).open("x", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")


def import_feedback(source_path, video_dir, output_dir):
    output_dir = Path(output_dir).resolve()
    if output_dir.exists():
        raise FileExistsError(f"Refusing to overwrite an existing review import: {output_dir}")
    source, keeper_rows, shooter_rows = read_source(source_path)
    videos = {clip: video_metadata(Path(video_dir) / f"{clip}.mp4") for clip in KEEPER_IDS}
    keeper = keeper_observations(source, keeper_rows, videos)
    shooter = shooter_observations(source, shooter_rows, videos)
    records = keeper + shooter
    for video in videos.values():
        if sha256(video["path"]) != video["sha256"]:
            raise ValueError(f"Source video changed during import: {video['video_id']}")
    if sha256(source["path"]) != source["sha256"]:
        raise ValueError("Source workbook changed during import")
    payloads = {
        "keeper_person_observations.json": {"schema_version": "pose-keeper-person-review-v1", "observations": keeper},
        "shooter_issue_ranges.json": {"schema_version": "pose-shooter-issue-review-v1", "observations": shooter},
        "regression_checklist.json": regression_checklist(records),
    }
    output_dir.mkdir(parents=True, exist_ok=False)
    for name, payload in payloads.items():
        write_json(output_dir / name, payload)
    manifest = {
        "schema_version": "pose-review-feedback-qc-v1",
        "imported_at_utc": datetime.now(timezone.utc).isoformat(),
        "importer": {"path": str(Path(__file__).resolve()), "sha256": sha256(__file__)},
        "source_workbook": source, "source_videos": list(videos.values()),
        "model_binding": binding(), "usage_constraints": usage_constraints(),
        "counts": {
            "keeper_person_observations": len(keeper), "shooter_issue_ranges": len(shooter),
            "source_videos": len(videos), "whole_clip_goalkeeper_misses": len(FULL_MISSES),
            "shooter_ranges_by_video": dict(Counter(row["source_video"]["video_id"] for row in shooter)),
            "pending_regression_items": len(records), "coordinate_training_labels": 0,
        },
        "qc_checks": {
            "source_schema_and_rows_validated": True, "all_ranges_within_source_video_metadata": True,
            "source_hashes_rechecked": True, "source_workbook_modified": False,
            "new_model_inference_performed": False, "legacy_predictions_modified": False,
            "model_output_bound": False, "model_regression_passed": None, "accuracy_computed": False,
        },
        "limitations": [
            "Video metadata and hashes identify review material; decoded content was not independently human-reviewed by this import.",
            "The exact reviewed model output is unavailable; every track ID remains binding_pending.",
            "Skeleton overlay observations do not establish missing CSV coordinates or corrected human joint coordinates.",
            "Left/right ambiguity is not a confirmed swap, striking-foot label, or contact-time label.",
            "Reported normal at 12-02 applies to the local person-association check, not whole-clip goalkeeper coverage.",
        ],
        "artifacts": {name: {"sha256": sha256(output_dir / name)} for name in payloads},
    }
    write_json(output_dir / "qc_manifest.json", manifest)
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--video-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/pose_review_feedback_2026-10-27")
    args = parser.parse_args()
    manifest = import_feedback(args.source, args.video_dir, args.output)
    print(json.dumps({"output": str(args.output.resolve()), "counts": manifest["counts"],
                      "model_binding": manifest["model_binding"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
