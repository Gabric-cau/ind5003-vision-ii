"""Separate label-free video inputs from coarse human contact-time references."""

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = ROOT / "outputs/core_contact_training_2026-10-26/data"
EXPECTED_GROUPS = {"col_eng": 4, "cro_den": 2, "esp_rus": 2, "rus_cro": 2}
CORE_SHEET = "\u8865\u680710\u6bb5"
REFERENCE_SHEET = "\u539f\u6807\u6ce839\u6bb5"


def sha256(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def probe_video(path):
    import cv2
    capture = cv2.VideoCapture(str(path))
    try:
        if not capture.isOpened():
            raise ValueError(f"Cannot open video: {path}")
        result = {"fps": float(capture.get(cv2.CAP_PROP_FPS)),
                  "frame_count": int(capture.get(cv2.CAP_PROP_FRAME_COUNT)),
                  "width": int(capture.get(cv2.CAP_PROP_FRAME_WIDTH)),
                  "height": int(capture.get(cv2.CAP_PROP_FRAME_HEIGHT))}
    finally:
        capture.release()
    if (not math.isfinite(result["fps"]) or result["fps"] <= 0
            or any(result[k] <= 0 for k in ("frame_count", "width", "height"))):
        raise ValueError(f"Invalid video metadata: {path}")
    result["duration_seconds"] = result["frame_count"] / result["fps"]
    return result


def finite_number(value, field):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f"Expected a finite numeric {field}")
    return float(value)


def build_documents(workbook_path, video_dir, *, probe=probe_video):
    from openpyxl import load_workbook
    workbook_path, video_dir = Path(workbook_path).resolve(), Path(video_dir).resolve()
    workbook_hash = sha256(workbook_path)
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    try:
        core = workbook[CORE_SHEET]
        reference = workbook[REFERENCE_SHEET]
        if core.cell(5, 7).value != "\u89e6\u7403\u65f6\u95f4(\u79d2)":
            raise ValueError("Unexpected contact-time column; source schema changed")
        source_hashes = {}
        for row in reference.iter_rows(min_row=6, values_only=True):
            if not row[0]:
                continue
            if row[0] in source_hashes:
                raise ValueError("Duplicate clip_id in source reference sheet")
            source_hashes[row[0]] = row[16]
        inputs, labels, checks = [], [], []
        seen_ids, seen_clips, seen_hashes = set(), set(), set()
        for row_number, row in enumerate(core.iter_rows(min_row=6, values_only=True), 6):
            if not row[0]:
                continue
            core_id, clip_id = row[0], row[1]
            if not isinstance(core_id, str) or not re.fullmatch(r"core_\d{2}", core_id):
                raise ValueError("Invalid core ID")
            if not isinstance(clip_id, str):
                raise ValueError("Missing source clip ID")
            match_group = "_".join(clip_id.split("_")[:2])
            if match_group not in EXPECTED_GROUPS:
                raise ValueError(f"Unexpected match group: {match_group}")
            if core_id in seen_ids or clip_id in seen_clips:
                raise ValueError("Duplicate event identity")
            seen_ids.add(core_id)
            seen_clips.add(clip_id)
            video = (video_dir / f"{core_id}.mp4").resolve()
            if not video.is_relative_to(video_dir):
                raise ValueError("Video escapes source directory")
            video_hash = sha256(video)
            if video_hash in seen_hashes:
                raise ValueError("Duplicate video bytes cannot be separate events")
            seen_hashes.add(video_hash)
            if source_hashes.get(clip_id) != video_hash:
                raise ValueError(f"Video hash differs from workbook current SHA256: {core_id}")
            media = probe(video)
            expected_fps = finite_number(row[5], "workbook fps")
            expected_duration = finite_number(row[3], "workbook duration")
            resolution = str(row[4]).replace("\u00d7", "x")
            if resolution != f"{media['width']}x{media['height']}":
                raise ValueError(f"Resolution mismatch: {core_id}")
            if abs(media["fps"] - expected_fps) > 1e-3:
                raise ValueError(f"FPS mismatch: {core_id}")
            if abs(media["duration_seconds"] - expected_duration) > 1 / media["fps"]:
                raise ValueError(f"Duration mismatch: {core_id}")
            contact = finite_number(row[6], "contact seconds")
            keeper = finite_number(row[13], "keeper movement seconds")
            camera_cut = finite_number(row[18], "first recorded camera cut seconds")
            if not (0 <= keeper <= contact < camera_cut <= media["duration_seconds"]):
                raise ValueError(f"Reference timing outside video/order bounds: {core_id}")
            if row[12] not in {"left", "center", "right", "unknown"}:
                raise ValueError(f"Unexpected keeper direction: {core_id}")
            if row[21] != "\u5df2\u5b8c\u6210\u590d\u6838":
                raise ValueError(f"Human review not marked completed: {core_id}")
            identity = {"core_id": core_id, "clip_id": clip_id, "match_group": match_group,
                        "video_sha256": video_hash}
            inputs.append({**identity, "video_path": str(video), **media})
            labels.append({
                **identity,
                "reference_contact_seconds": contact,
                "reference_keeper_movement_seconds": keeper,
                "reference_keeper_movement_direction": row[12],
                "reference_first_recorded_camera_cut_seconds": camera_cut,
                "reference_precision": "coarse_human_viewing_time_not_frame_exact",
                "reference_uncertainty_seconds": None,
                "reference_provenance": "human_filled_workbook_as_delivered",
                "source_workbook_sha256": workbook_hash,
                "source_sheet": CORE_SHEET,
                "source_row": row_number,
                "source_column": 7,
                "source_cell": f"G{row_number}",
                "keeper_source_cells": [f"M{row_number}", f"N{row_number}"],
                "review_status_as_recorded": row[21],
                "reviewer_identity_authenticated": False,
                "independent_second_reviewer_consensus": False,
                "old_reference_frames_used": False,
            })
            checks.append({"core_id": core_id, "current_video_sha256_matches_workbook": True,
                           "fps_resolution_duration_match_workbook": True,
                           "timing_bounds_valid": True,
                           "keeper_and_contact_same_recorded_time": keeper == contact,
                           "uncertainty_notes_empty": row[20] is None})
    finally:
        workbook.close()
    groups = Counter(row["match_group"] for row in inputs)
    if dict(groups) != EXPECTED_GROUPS or len(inputs) != 10:
        raise ValueError(f"Expected four match groups with 4/2/2/2 events, found {dict(groups)}")
    if seen_ids != {f"core_{i:02d}" for i in range(1, 11)}:
        raise ValueError("Expected core_01 through core_10")
    folds = [{"held_out_match": group,
              "train_core_ids": [row["core_id"] for row in inputs if row["match_group"] != group],
              "validation_core_ids": [row["core_id"] for row in inputs if row["match_group"] == group]}
             for group in EXPECTED_GROUPS]
    inputs_document = {
        "schema_version": "core-contact-inputs-v1",
        "samples": inputs,
        "folds": folds,
        "reference_labels_included": False,
        "all_views_and_frames_share_event_and_match_group": True,
    }
    labels_document = {
        "schema_version": "core-contact-references-v1",
        "source_workbook": str(workbook_path),
        "source_workbook_sha256": workbook_hash,
        "samples": labels,
        "training_target": "reference_contact_seconds",
        "other_reference_fields_are_audit_only": True,
        "old_frame_offsets_transferred": False,
    }
    audit_document = {
        "schema_version": "core-contact-preparation-audit-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "source_workbook": str(workbook_path),
        "source_workbook_sha256": workbook_hash,
        "selected_events": len(inputs),
        "match_group_counts": dict(groups),
        "validation_protocol": "four_fold_leave_one_match_out_development_evaluation",
        "folds": folds,
        "source_events_are_subset_of_existing_39": True,
        "counted_as_new_independent_events": False,
        "final_untouched_test_evaluated": False,
        "external_world_cup_eight_accessed": False,
        "old_frame_offsets_transferred": False,
        "media_transformed": False,
        "source_workbook_modified": False,
        "checks": checks,
        "limitations": [
            "Contact references are human viewing times, not frame-exact physical contact ground truth.",
            "No per-reference uncertainty interval was supplied; none is inferred.",
            "core_01 and core_03 have identical recorded keeper/contact times.",
            "Ten selected events from four previously used development matches are not an untouched final test.",
            "Old reference video frames are not mapped by assuming a constant offset.",
            "Keeper timing and manually recorded camera cuts must not enter input feature extraction.",
        ],
    }
    return inputs_document, labels_document, audit_document


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--videos", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    documents = build_documents(args.workbook, args.videos)
    args.output.mkdir(parents=True, exist_ok=True)
    for name, document in zip(("inputs", "labels", "audit"), documents):
        (args.output / f"{name}.json").write_text(
            json.dumps(document, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "events": len(documents[0]["samples"]),
                      "match_group_counts": documents[2]["match_group_counts"]}, indent=2))


if __name__ == "__main__":
    main()
