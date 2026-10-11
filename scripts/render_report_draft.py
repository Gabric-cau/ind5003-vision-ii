"""Render the six-section review draft to Times New Roman PDF and DOCX."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt
from markdown_it.token import Token
from PIL import Image as PillowImage
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
FONT_FILES = {
    "TNR": "Times New Roman.ttf",
    "TNR-Bold": "Times New Roman Bold.ttf",
    "TNR-Italic": "Times New Roman Italic.ttf",
    "TNR-BoldItalic": "Times New Roman Bold Italic.ttf",
}
MARGIN = 0.8 * inch
PAGE_WIDTH, PAGE_HEIGHT = A4
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN
CONTENT_HEIGHT = PAGE_HEIGHT - 2 * MARGIN


@dataclass(frozen=True)
class Block:
    kind: str
    children: list[Token]
    level: int = 0


def plain_text(children: list[Token]) -> str:
    return "".join(" " if token.type in {"softbreak", "hardbreak"} else token.content
                   for token in children if token.type in {"text", "code_inline", "softbreak", "hardbreak"})


def parse_report(source: Path) -> list[Block]:
    tokens = MarkdownIt("commonmark").parse(source.read_text(encoding="utf-8"))
    blocks = []
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type not in {"heading_open", "paragraph_open"}:
            raise ValueError(f"Unsupported report block: {token.type}; nothing was silently omitted")
        if index + 2 >= len(tokens) or tokens[index + 1].type != "inline":
            raise ValueError("Expected structured Markdown inline content")
        children = tokens[index + 1].children or []
        expected_close = "heading_close" if token.type == "heading_open" else "paragraph_close"
        if tokens[index + 2].type != expected_close:
            raise ValueError("Malformed Markdown block")
        images = [child for child in children if child.type == "image"]
        if images:
            if len(children) != 1 or token.type != "paragraph_open":
                raise ValueError("Place each report figure in its own Markdown paragraph")
            blocks.append(Block("figure", children))
        elif token.type == "heading_open":
            level = int(token.tag[1:])
            if level not in {1, 2}:
                raise ValueError("The report renderer supports its title and six section headings")
            blocks.append(Block("heading", children, level))
        else:
            blocks.append(Block("paragraph", children))
        index += 3
    if sum(block.kind == "heading" and block.level == 1 for block in blocks) != 1:
        raise ValueError("Expected exactly one report title")
    if sum(block.kind == "heading" and block.level == 2 for block in blocks) != 6:
        raise ValueError("Expected exactly six course report sections")
    if "review draft" not in source.read_text(encoding="utf-8").lower():
        raise ValueError("The source must retain its review draft status")
    return blocks


def figure_path(source: Path, token: Token) -> Path:
    target = token.attrGet("src") or ""
    if "://" in target:
        raise ValueError("Report figures must be available local files")
    path = (source.parent / target).resolve()
    if not path.is_file():
        raise ValueError(f"Missing figure: {path}")
    return path


def figure_dimensions(path: Path) -> tuple[float, float]:
    with PillowImage.open(path) as image:
        width, height = image.size
    scale = min(CONTENT_WIDTH / width, CONTENT_HEIGHT * 0.52 / height)
    return width * scale, height * scale


def pdf_inline(children: list[Token]) -> str:
    output = []
    for token in children:
        if token.type == "text":
            output.append(html.escape(token.content))
        elif token.type == "code_inline":
            output.append(f'<font color="#444444">{html.escape(token.content)}</font>')
        elif token.type in {"softbreak", "hardbreak"}:
            output.append(" " if token.type == "softbreak" else "<br/>")
        elif token.type == "strong_open":
            output.append("<b>")
        elif token.type == "strong_close":
            output.append("</b>")
        elif token.type == "em_open":
            output.append("<i>")
        elif token.type == "em_close":
            output.append("</i>")
        elif token.type == "link_open":
            target = html.escape(token.attrGet("href") or "", quote=True)
            output.append(f'<link href="{target}" color="#164c78">')
        elif token.type == "link_close":
            output.append("</link>")
        else:
            raise ValueError(f"Unsupported report inline token: {token.type}")
    return "".join(output)


def register_fonts(directory: Path) -> None:
    for name, filename in FONT_FILES.items():
        path = directory / filename
        if not path.is_file():
            raise ValueError(f"Required Times New Roman font is missing: {path}")
        pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("TNR", normal="TNR", bold="TNR-Bold",
                                 italic="TNR-Italic", boldItalic="TNR-BoldItalic")


def page_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("TNR", 10)
    canvas.setFillColor(colors.HexColor("#555555"))
    canvas.drawString(MARGIN, MARGIN * 0.45, "IND5003 Vision II | Review draft")
    canvas.drawRightString(PAGE_WIDTH - MARGIN, MARGIN * 0.45, str(document.page))
    canvas.restoreState()


def write_pdf(source: Path, destination: Path, blocks: list[Block], font_dir: Path) -> int:
    register_fonts(font_dir)
    styles = {
        "title": ParagraphStyle("ReportTitle", fontName="TNR-Bold", fontSize=16,
                                leading=20, spaceAfter=12, keepWithNext=True),
        "heading": ParagraphStyle("ReportSection", fontName="TNR-Bold", fontSize=13,
                                  leading=16, spaceBefore=12, spaceAfter=7, keepWithNext=True),
        "body": ParagraphStyle("ReportBody", fontName="TNR", fontSize=11, leading=14,
                               alignment=TA_JUSTIFY, spaceAfter=7, allowWidows=0, allowOrphans=0),
        "caption": ParagraphStyle("FigureCaption", fontName="TNR-Italic", fontSize=10,
                                  leading=12, alignment=TA_CENTER, spaceBefore=5, spaceAfter=9),
    }
    title = next(plain_text(block.children) for block in blocks
                 if block.kind == "heading" and block.level == 1)
    document = SimpleDocTemplate(str(destination), pagesize=A4, leftMargin=MARGIN,
                                 rightMargin=MARGIN, topMargin=MARGIN, bottomMargin=MARGIN,
                                 title=title, author="", subject="IND5003 report review draft")
    story = []
    for block in blocks:
        if block.kind == "figure":
            token = block.children[0]
            path = figure_path(source, token)
            width, height = figure_dimensions(path)
            story.append(KeepTogether([Spacer(1, 5), Image(str(path), width=width, height=height),
                                       Paragraph(html.escape(token.content), styles["caption"])]))
        else:
            style_name = "title" if block.level == 1 else "heading" if block.kind == "heading" else "body"
            story.append(Paragraph(pdf_inline(block.children), styles[style_name]))
    document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    pages = len(PdfReader(destination).pages)
    if pages > 25:
        raise ValueError(f"The report exceeds the 25-page main-text limit: {pages} pages")
    return pages


def set_font(run, size: float = 11, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")


def add_hyperlink(paragraph, text: str, target: str, *, size: float, bold: bool, italic: bool) -> None:
    relationship = paragraph.part.relate_to(
        target, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), "Times New Roman")
    properties.append(fonts)
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(round(size * 2)))
    properties.append(font_size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "164C78")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    content = OxmlElement("w:t")
    content.set(qn("xml:space"), "preserve")
    content.text = text
    run.append(content)
    link.append(run)
    paragraph._p.append(link)


def docx_inline(paragraph, children: list[Token], *, size: float = 11, heading: bool = False) -> None:
    bold_depth = italic_depth = 0
    target = None
    for token in children:
        if token.type == "strong_open":
            bold_depth += 1
        elif token.type == "strong_close":
            bold_depth -= 1
        elif token.type == "em_open":
            italic_depth += 1
        elif token.type == "em_close":
            italic_depth -= 1
        elif token.type == "link_open":
            target = token.attrGet("href")
        elif token.type == "link_close":
            target = None
        elif token.type in {"text", "code_inline", "softbreak", "hardbreak"}:
            content = " " if token.type == "softbreak" else "\n" if token.type == "hardbreak" else token.content
            bold, italic = bool(bold_depth or heading), bool(italic_depth)
            if target:
                add_hyperlink(paragraph, content, target, size=size, bold=bold, italic=italic)
            else:
                set_font(paragraph.add_run(content), size, bold, italic)
        else:
            raise ValueError(f"Unsupported report inline token: {token.type}")


def write_docx(source: Path, destination: Path, blocks: list[Block]) -> None:
    document = Document()
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.title = next(plain_text(block.children) for block in blocks
                                          if block.kind == "heading" and block.level == 1)
    document.core_properties.subject = "IND5003 report review draft"
    section = document.sections[0]
    section.page_width, section.page_height = Inches(PAGE_WIDTH / inch), Inches(PAGE_HEIGHT / inch)
    section.top_margin = section.bottom_margin = Inches(MARGIN / inch)
    section.left_margin = section.right_margin = Inches(MARGIN / inch)
    for name in ("Normal", "Title", "Heading 1", "Caption", "Footer"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11 if name == "Normal" else 10)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    normal = document.styles["Normal"].paragraph_format
    normal.line_spacing = 1.15
    normal.space_after = Pt(7)
    normal.widow_control = True
    footer = section.footer.paragraphs[0]
    set_font(footer.add_run("IND5003 Vision II | Review draft | Page "), 10)
    page = footer.add_run("1")
    set_font(page, 10)
    footer._p.remove(page._r)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field.append(page._r)
    footer._p.append(field)
    for block in blocks:
        if block.kind == "figure":
            token = block.children[0]
            path = figure_path(source, token)
            width, height = figure_dimensions(path)
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Inches(width / inch), height=Inches(height / inch))
            caption = document.add_paragraph(style="Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.keep_together = True
            set_font(caption.add_run(token.content), 10, italic=True)
        elif block.kind == "heading":
            paragraph = document.add_paragraph(style="Title" if block.level == 1 else "Heading 1")
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(0 if block.level == 1 else 12)
            paragraph.paragraph_format.space_after = Pt(12 if block.level == 1 else 7)
            docx_inline(paragraph, block.children, size=16 if block.level == 1 else 13, heading=True)
        else:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            docx_inline(paragraph, block.children)
    document.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "docs/REPORT_REVIEW_DRAFT.md")
    parser.add_argument("--output-dir", type=Path, help="Defaults to the Markdown source directory")
    parser.add_argument("--font-dir", type=Path, default=Path("/System/Library/Fonts/Supplemental"),
                        help="Directory containing all four Times New Roman TTF files")
    args = parser.parse_args()
    source = args.source.resolve()
    blocks = parse_report(source)
    output_dir = (args.output_dir or source.parent).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path, docx_path = output_dir / f"{source.stem}.pdf", output_dir / f"{source.stem}.docx"
    pages = write_pdf(source, pdf_path, blocks, args.font_dir)
    write_docx(source, docx_path, blocks)
    print(json.dumps({"source": str(source), "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                      "pdf": str(pdf_path), "docx": str(docx_path), "pdf_pages": pages,
                      "body_font": "Times New Roman", "body_font_pt": 11,
                      "minimum_text_font_pt": 10, "status": "review_draft"}, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
