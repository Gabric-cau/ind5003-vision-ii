import cv2
import numpy as np
import openpyxl
import pytest

from scripts.prepare_ball_frame_annotations import (
    FRAME_HEADERS, FRAME_SHEET, NOTES_SHEET, SUMMARY_SHEET, SPLITS,
    normalize_rows, prepare, read_json, sha256,
)


@pytest.fixture
def source_bundle(tmp_path):
    roots = {"core_root": tmp_path / "core", "match_root": tmp_path / "match"}
    for root in roots.values():
        root.mkdir()
    workbook = openpyxl.Workbook()
    frames = workbook.active
    frames.title = FRAME_SHEET
    frames.append(FRAME_HEADERS)
    summary = workbook.create_sheet(SUMMARY_SHEET)
    summary.append(["id", "description", "contact", "last", "position", "outcome", "cut", "note"])
    notes = workbook.create_sheet(NOTES_SHEET)
    notes.append(["Fixture preserves source rows; coordinates are local pixels."])
    for clip in SPLITS:
        fps = 25 if clip == "match_01" else 30
        path = roots["match_root" if clip == "match_01" else "core_root"] / f"{clip}.mp4"
        writer = cv2.VideoWriter(str(path), cv2.VideoWriter_fourcc(*"mp4v"), fps, (64, 48))
        assert writer.isOpened()
        for frame in range(8):
            image = np.full((48, 64, 3), frame * 20, dtype=np.uint8)
            cv2.rectangle(image, (20, 12), (28, 20), (255, 255, 255), -1)
            writer.write(image)
        writer.release()
        frames.append([None, None, clip, round(2 / fps, 3), 2, 24, 16, 20, 12, 28, 20,
                       "\u6e05\u6670", "refined", "\u89e6\u7403", None])
        summary.append([clip, f"64x48 {fps}fps", "1\u21922\u4e4b\u95f4", None, None, None, None, None])
    extra = [
        (3, "\u90e8\u5206\u53ef\u89c1", [24, 16, 20, 12, 28, 20], None),
        (4, "\u6e05\u6670", [24, 16, None, None, None, None], "\u4eba\u5de5\u63d2\u503c\u4f30\u8ba1"),
        (5, "\u770b\u4e0d\u6e05", [None] * 6, "\u51fa\u753b\u9762"),
        (6, "\u753b\u9762\u5916\uff08\u955c\u5934\u5207\u8d70\uff09", [None] * 6, None),
        (7, "\u6e05\u6670", [24, 16, None, None, None, None], "\u534a\u622a\u51fa\u753b"),
    ]
    for frame, visibility, geometry, note in extra:
        frames.append([None, None, "core_03", round(frame / 30, 3), frame, *geometry,
                       visibility, "manual", "after_contact", note])
    source = tmp_path / "source.xlsx"
    workbook.save(source)
    workbook.close()
    return source, roots


def test_preserves_raw_source_intervals_and_exports_only_eligible_frames(source_bundle, tmp_path):
    source, roots = source_bundle
    before = sha256(source)
    output = tmp_path / "output"
    result = prepare(source, output, **roots, export_yolo=True)
    assert sha256(source) == before
    assert result["rows"] == 9
    assert result["detection_train_eligible"] == result["yolo_exported_images"] == 4
    assert result["eligible_by_split"] == {"train": 3, "validation": 1}
    assert result["negative_samples"] == 0
    document = read_json(output / "annotations.json")
    rows = document["annotations"]
    assert rows[0]["source"]["row_1_based"] == 2
    assert rows[0]["source"]["sheet"] == FRAME_SHEET
    assert rows[0]["source"]["workbook_sha256"] == before
    assert rows[0]["source"]["raw_values"][5:11] == [24, 16, 20, 12, 28, 20]
    assert document["contact_intervals"][0]["frame_interval_zero_based"] == [1, 2]
    assert document["contact_intervals"][0]["is_single_frame_ground_truth"] is False
    assert rows[4]["partial_visibility_candidate"] is True
    assert rows[5]["interpolated"] is True and rows[5]["bbox_xyxy_px"] is None
    assert rows[6]["visibility"] == "out_of_view"
    assert rows[7]["visibility"] == "shot_cut"
    assert rows[8]["truncated"] is True and rows[8]["bbox_xyxy_px"] is None
    assert all(not row["negative_sample_eligible"] for row in rows)
    manifest = read_json(output / "manifest.json")
    assert manifest["unseen_data_claim"] is False
    assert manifest["split_policy"]["frame_random_split"] is False
    assert set(manifest["videos"]) == set(SPLITS)
    for video in manifest["videos"].values():
        assert video["decoded_frame_count"] == video["reported_frame_count"] == 8
        assert sha256(video["path"]) == video["sha256"]
    for exported in manifest["exports"]:
        image = cv2.imread(str(output / exported["image"]))
        assert image.shape[:2] == (48, 64)
        assert image.mean() == pytest.approx(40, abs=12)
        label = (output / exported["label"]).read_text().split()
        assert label[0] == "0"
        assert [float(x) for x in label[1:]] == pytest.approx([24 / 64, 16 / 48, 8 / 64, 8 / 48])
        assert exported["split"] == ("validation" if exported["annotation_id"].startswith("match") else "train")
    assert "0: ball" in (output / "yolo/dataset.yaml").read_text()
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        prepare(source, output, **roots)


@pytest.mark.parametrize("column,value,error", [
    (5, 99, "Frame outside video"),
    (4, 1.0, "Frame/time mismatch"),
    (8, 30, "Invalid bounding box geometry"),
    (6, 40, "Center outside bounding box"),
    (9, None, "Incomplete bounding box"),
    (3, "core_05", "not allowlisted"),
])
def test_invalid_annotations_fail_without_output(source_bundle, tmp_path, column, value, error):
    source, roots = source_bundle
    workbook = openpyxl.load_workbook(source)
    workbook[FRAME_SHEET].cell(2, column).value = value
    workbook.save(source)
    workbook.close()
    before = sha256(source)
    with pytest.raises(ValueError, match=error):
        prepare(source, tmp_path / "bad", **roots)
    assert not (tmp_path / "bad").exists()
    assert sha256(source) == before


def test_duplicate_frame_rejected(source_bundle, tmp_path):
    source, roots = source_bundle
    workbook = openpyxl.load_workbook(source)
    sheet = workbook[FRAME_SHEET]
    sheet.append([cell.value for cell in sheet[2]])
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="Duplicate video/frame"):
        prepare(source, tmp_path / "duplicate", **roots)


def test_declared_video_metadata_mismatch_rejected(source_bundle, tmp_path):
    source, roots = source_bundle
    workbook = openpyxl.load_workbook(source)
    workbook[SUMMARY_SHEET].cell(2, 2).value = "1280x720 30fps"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ValueError, match="dimensions/fps"):
        prepare(source, tmp_path / "mismatch", **roots)


def test_interpolation_and_truncation_excluded_even_with_box():
    videos = {clip: {"width": 64, "height": 48, "fps": 30, "reported_frame_count": 8} for clip in SPLITS}
    source_rows = [{"raw_values": FRAME_HEADERS}]
    for clip in SPLITS:
        note = "\u63d2\u503c" if clip == "core_01" else "\u534a\u622a\u51fa\u753b" if clip == "core_02" else None
        source_rows.append({"raw_values": [None, None, clip, 0, 0, 24, 16, 20, 12, 28, 20,
                                          "\u6e05\u6670", "refined", "before", note]})
    rows = normalize_rows(source_rows, videos)
    assert [row["detection_train_eligible"] for row in rows] == [False, False, True, True]
