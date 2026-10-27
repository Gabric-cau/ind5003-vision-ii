import json

import cv2
import numpy as np
import openpyxl
import pytest

from scripts.prepare_ball_frame_annotations import FRAME_HEADERS, FRAME_SHEET, SUMMARY_SHEET, NOTES_SHEET, sha256
from scripts.prepare_ball_frame_evaluation import VIDEO_IDS, assign_duplicate_groups, normalize_rows, prepare


@pytest.fixture
def reference_bundle(tmp_path):
    root = tmp_path / "playback"
    root.mkdir()
    workbook = openpyxl.Workbook()
    frames = workbook.active
    frames.title = FRAME_SHEET
    frames.append(FRAME_HEADERS)
    summary = workbook.create_sheet(SUMMARY_SHEET)
    summary.append(["id", "description", "contact", "last", "position", "outcome", "cut", "note"])
    workbook.create_sheet(NOTES_SHEET).append(["Evaluation only"])
    for clip in VIDEO_IDS:
        writer = cv2.VideoWriter(str(root / f"{clip}.mp4"), cv2.VideoWriter_fourcc(*"mp4v"), 30, (64, 48))
        assert writer.isOpened()
        for frame in range(10):
            image = np.full((48, 64, 3), frame * 15, dtype=np.uint8)
            cv2.rectangle(image, (20, 12), (28, 20), (255, 255, 255), -1)
            writer.write(image)
        writer.release()
        frames.append([None, None, clip, round(2 / 30, 3), 2, 24, 16, 20, 12, 28, 20,
                       "\u6e05\u6670", "refined", "before", None])
        summary.append([clip, "64x48 30fps", "1\u21922\u4e4b\u95f4", None, None, None, None, None])
    extra = [
        (3, "\u6e05\u6670", [24, 16, 20, 12, 28, 20], "\u4e0e\u4e0a\u4e00\u5e27\u753b\u9762\u91cd\u590d"),
        (4, "\u6e05\u6670", [24, 16, 20, 12, 28, 20], "\u4eba\u5de5\u63d2\u503c"),
        (5, "\u90e8\u5206\u53ef\u89c1", [24, 16, None, None, None, None], "occluded"),
        (6, "\u6e05\u6670", [24, 16, None, None, None, None], "\u540c\u4e0a\uff08\u4e0e\u4e0a\u4e00\u5e27\u753b\u9762\u91cd\u590d\uff09"),
        (7, "\u6e05\u6670", [24, 16, 20, 12, 28, 20], "\u534a\u622a\u5728\u753b\u9762\u8fb9\u7f18\u5916"),
        (8, "\u770b\u4e0d\u6e05", [None] * 6, "\u51fa\u753b\u9762"),
    ]
    for frame, visibility, geometry, note in extra:
        frames.append([None, None, "core_09", round(frame / 30, 3), frame, *geometry, visibility, "manual", "after", note])
    source = tmp_path / "reference.xlsx"
    workbook.save(source)
    workbook.close()
    return source, root


def test_evaluation_export_preserves_source_and_never_creates_training_data(reference_bundle, tmp_path):
    source, root = reference_bundle
    before = sha256(source)
    output = tmp_path / "evaluation"
    result = prepare(source, output, core_root=root)
    assert sha256(source) == before
    assert result["rows"] == 9
    assert result["evaluation_frames"] == 4
    assert result["deduplicated_evaluation_frames"] == 3
    assert result["negative_samples"] == 0
    assert not (output / "yolo/images/train").exists()
    assert not (output / "yolo/dataset.yaml").exists()
    document = json.loads((output / "annotations.json").read_text())
    rows = document["annotations"]
    assert document["evaluation_only"] and not document["training_permitted"]
    assert all(not row["detection_train_eligible"] for row in rows)
    assert rows[0]["source"]["workbook_sha256"] == before
    assert rows[0]["source"]["raw_values"][5:11] == [24, 16, 20, 12, 28, 20]
    assert rows[3]["declared_repeat_of_previous_frame"]
    assert rows[3]["deduplicated_representative_id"] == "core_09_f000002"
    assert not rows[4]["detection_evaluation_eligible"]
    assert rows[6]["visibility"] == "partial"
    assert rows[6]["quality_flags"] == ["copied_previous_frame_has_lower_visibility"]
    assert rows[7]["truncated"] and not rows[7]["detection_evaluation_eligible"]
    assert rows[8]["visibility"] == "out_of_view"
    manifest = json.loads((output / "manifest.json").read_text())
    for export in manifest["exports"]:
        assert export["split"] == "test"
        assert sha256(output / export["image"]) == export["image_sha256"]
        assert sha256(output / export["label"]) == export["label_sha256"]
        assert export["normalized"] == pytest.approx([24 / 64, 16 / 48, 8 / 64, 8 / 48])
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        prepare(source, output, core_root=root)


@pytest.mark.parametrize("column,value,error", [
    (5, 99, "Frame outside video"),
    (4, 1.0, "Frame/time mismatch"),
    (8, 30, "Invalid bounding box geometry"),
    (9, None, "Incomplete bounding box"),
    (3, "core_01", "evaluation allowlist"),
])
def test_invalid_reference_does_not_leave_output(reference_bundle, tmp_path, column, value, error):
    source, root = reference_bundle
    workbook = openpyxl.load_workbook(source)
    workbook[FRAME_SHEET].cell(2, column).value = value
    workbook.save(source)
    workbook.close()
    output = tmp_path / "invalid"
    with pytest.raises(ValueError, match=error):
        prepare(source, output, core_root=root)
    assert not output.exists()


def test_reentry_does_not_claim_same_ball_identity():
    videos = {clip: {"width": 64, "height": 48, "fps": 30, "reported_frame_count": 200} for clip in VIDEO_IDS}
    sources = [{"raw_values": FRAME_HEADERS}]
    for clip in VIDEO_IDS:
        frame = 111 if clip == "core_09" else 2
        sources.append({"raw_values": [None, None, clip, round(frame / 30, 3), frame, 24, 16,
                                       20, 12, 28, 20, "\u6e05\u6670", "refined", "after", None]})
    rows = normalize_rows(sources, videos)
    assert rows[-1]["detection_evaluation_eligible"]
    assert rows[-1]["identity_uncertain_after_reentry"]
    assert not rows[-1]["continuous_shot_identity_reference"]


def test_exact_pixel_duplicate_group_uses_first_eligible_frame_only():
    rows = [{"id": f"core_05_f{frame:06d}", "video_id": "core_05", "frame_zero_based": frame,
             "decoded_pixels_sha256": digest, "declared_repeat_of_previous_frame": False,
             "detection_evaluation_eligible": eligible}
            for frame, digest, eligible in [(1, "same", False), (3, "same", True), (4, "same", True), (7, "other", True)]]
    assign_duplicate_groups(rows)
    assert [row["deduplicated_primary"] for row in rows] == [False, True, False, True]
    assert rows[2]["deduplicated_representative_id"] == rows[1]["id"]
