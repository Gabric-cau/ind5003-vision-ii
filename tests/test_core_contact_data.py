import hashlib

from openpyxl import Workbook, load_workbook
import pytest

from scripts.prepare_core_contact_training import CORE_SHEET, REFERENCE_SHEET, build_documents


def fixture(tmp_path):
    videos = tmp_path / "videos"
    videos.mkdir()
    workbook = Workbook()
    core = workbook.active
    core.title = CORE_SHEET
    reference = workbook.create_sheet(REFERENCE_SHEET)
    core.cell(5, 7, "\u89e6\u7403\u65f6\u95f4(\u79d2)")
    groups = ["col_eng"] * 4 + ["cro_den"] * 2 + ["esp_rus"] * 2 + ["rus_cro"] * 2
    for index, group in enumerate(groups, 1):
        row = index + 5
        core_id = f"core_{index:02d}"
        clip_id = f"{group}_p{index:02d}_player"
        content = core_id.encode()
        (videos / f"{core_id}.mp4").write_bytes(content)
        values = {1: core_id, 2: clip_id, 4: 10, 5: "1280\u00d7720", 6: 30,
                  7: 2.5, 13: "left", 14: 2.3, 19: 4.5,
                  22: "\u5df2\u5b8c\u6210\u590d\u6838"}
        for column, value in values.items():
            core.cell(row, column, value)
        reference.cell(row, 1, clip_id)
        reference.cell(row, 17, hashlib.sha256(content).hexdigest())
    path = tmp_path / "references.xlsx"
    workbook.save(path)
    return path, videos


def probe(_):
    return {"fps": 30.0, "frame_count": 300, "width": 1280, "height": 720,
            "duration_seconds": 10.0}


def test_inputs_are_label_free_and_match_folds_never_share_events(tmp_path):
    path, videos = fixture(tmp_path)
    inputs, labels, audit = build_documents(path, videos, probe=probe)
    assert len(inputs["samples"]) == len(labels["samples"]) == 10
    for row in inputs["samples"]:
        assert set(row) == {"core_id", "clip_id", "match_group", "video_sha256", "video_path",
                            "fps", "frame_count", "width", "height", "duration_seconds"}
    for fold in inputs["folds"]:
        assert not set(fold["train_core_ids"]) & set(fold["validation_core_ids"])
        group_map = {row["core_id"]: row["match_group"] for row in inputs["samples"]}
        assert {group_map[cid] for cid in fold["validation_core_ids"]} == {fold["held_out_match"]}
        assert fold["held_out_match"] not in {group_map[cid] for cid in fold["train_core_ids"]}
    assert sorted(len(fold["validation_core_ids"]) for fold in inputs["folds"]) == [2, 2, 2, 4]
    assert labels["samples"][0]["source_cell"] == "G6"
    assert labels["samples"][0]["reference_uncertainty_seconds"] is None
    assert audit["external_world_cup_eight_accessed"] is False


def test_changed_media_bytes_fail_before_training(tmp_path):
    path, videos = fixture(tmp_path)
    (videos / "core_01.mp4").write_bytes(b"modified")
    with pytest.raises(ValueError, match="Video hash differs"):
        build_documents(path, videos, probe=probe)


@pytest.mark.parametrize("column,value,message", [
    (7, 11, "timing outside"), (7, "2.5", "finite numeric"),
    (14, 3, "timing outside"), (6, 25, "FPS mismatch"),
    (5, "1920x1080", "Resolution mismatch"), (22, "pending", "not marked completed"),
])
def test_invalid_timing_media_or_review_is_rejected(tmp_path, column, value, message):
    path, videos = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook[CORE_SHEET].cell(6, column, value)
    workbook.save(path)
    with pytest.raises(ValueError, match=message):
        build_documents(path, videos, probe=probe)


def test_duplicate_event_and_missing_match_are_rejected(tmp_path):
    path, videos = fixture(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook[CORE_SHEET]
    sheet.cell(7, 2, sheet.cell(6, 2).value)
    workbook.save(path)
    with pytest.raises(ValueError, match="Duplicate event identity"):
        build_documents(path, videos, probe=probe)


def test_same_keeper_contact_time_is_retained_as_coarse_reference(tmp_path):
    path, videos = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook[CORE_SHEET].cell(6, 14, 2.5)
    workbook.save(path)
    _, labels, audit = build_documents(path, videos, probe=probe)
    assert labels["samples"][0]["reference_contact_seconds"] == 2.5
    assert audit["checks"][0]["keeper_and_contact_same_recorded_time"] is True
    assert labels["old_frame_offsets_transferred"] is False
