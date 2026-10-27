import json

import openpyxl
import pytest

from scripts import import_pose_review_feedback as importer


@pytest.fixture
def source_workbook(tmp_path):
    workbook = openpyxl.Workbook()
    keeper = workbook.active
    keeper.title = importer.KEEPER_SHEET
    keeper.cell(2, 1, "YOLO11m-pose 人工复核反馈")
    keeper.cell(4, 1, "帧号从0开始。")
    shooter = workbook.create_sheet(importer.SHOOTER_SHEET)
    shooter.cell(3, 1, "帧范围包含两端，帧号从0开始。")
    for sheet in (keeper, shooter):
        for column, header in enumerate(importer.HEADERS[sheet.title], start=1):
            sheet.cell(6, column, header)
    for number, clip in enumerate(importer.KEEPER_IDS, start=7):
        for column, value in enumerate((clip, *importer.KEEPER_CONTRACT[clip], "原文归类", "原文歧义", "原文处理"), start=1):
            keeper.cell(number, column, value)
    keeper.cell(18, 1, "用户原话保留")
    for number, row in enumerate(importer.SHOOTER_ROWS, start=7):
        for column, value in enumerate((*row, "左右腿", "骨架丢失、左右混淆", "原文原因", "原文范围", None), start=1):
            shooter.cell(number, column, value)
    path = tmp_path / "source.xlsx"
    workbook.save(path)
    workbook.close()
    return path


@pytest.fixture
def videos():
    lengths = dict(zip(importer.KEEPER_IDS, (80, 137, 207, 134, 141, 156, 126, 129)))
    return {clip: {"video_id": clip, "sha256": "a" * 64, "frame_count": length} for clip, length in lengths.items()}


def records(path, videos):
    source, keeper_rows, shooter_rows = importer.read_source(path)
    return source, importer.keeper_observations(source, keeper_rows, videos), importer.shooter_observations(source, shooter_rows, videos)


def test_twelve_exact_ranges_preserve_source_and_are_not_training_labels(source_workbook, videos):
    original_hash = importer.sha256(source_workbook)
    source, keeper, shooter = records(source_workbook, videos)
    assert len(keeper) == 8
    assert len(shooter) == 12
    assert importer.sha256(source_workbook) == original_hash == source["sha256"]
    assert source["sheets"][0]["nonempty_rows"][-1] == {"row": 18, "cells": ["用户原话保留", None, None, None, None, None, None]}
    for offset, (record, expected) in enumerate(zip(shooter, importer.SHOOTER_ROWS), start=7):
        clip, model_id, start, end, recovery = expected
        assert record["source_video"]["video_id"] == clip
        assert record["source_video"]["sha256"] == "a" * 64
        assert record["subject_model_id"] == model_id
        assert record["issue_frame_range"] == {"start": start, "end": end, "inclusive": True, "index_base": 0}
        assert record["reported_recovery_frame"] == recovery
        assert record["source"]["row"] == offset
        assert record["source"]["sheet"] == importer.SHOOTER_SHEET
        assert record["source"]["workbook_sha256"] == original_hash
        assert record["reported_observation"] == record["source"]["raw_cells"]["用户目视观察"]
        assert record["usage_constraints"]["training_label_eligible"] is False
        assert record["coordinate_labels"] is None
        assert record["striking_foot_label"] is None
        assert record["ball_contact_frame_label"] is None
        assert record["specific_missing_joint_indices"] is None
        assert record["actual_left_right_swap_confirmed"] is False
        assert record["csv_coordinate_missing_confirmed"] is False
        assert record["corrected_coordinates_available"] is False
        assert record["model_binding"]["reviewed_model"] == "YOLO11m-pose"
        assert record["model_binding"]["status"] == "binding_pending"
        assert record["model_binding"]["reviewed_inference_output_sha256"] is None
        assert record["model_binding"]["legacy_output_id_mapping_allowed"] is False
    assert shooter[0]["range_basis"] == "start_and_recovery_implied_interval"
    assert [shooter[index]["issue_frame_range"]["start"] for index in (1, 2)] == [60, 65]


def test_unknown_fragmentation_and_local_pass_are_kept_distinct(source_workbook, videos):
    _, keeper, shooter = records(source_workbook, videos)
    indexed = {row["source_video"]["video_id"]: row for row in keeper}
    unknown = indexed["07-08"]
    assert unknown["subject_model_ids"] == [2]
    assert unknown["frame_range_states"] == ["unknown", "reported_normal"]
    assert unknown["review_frame_ranges"][0] == {"start": 67, "end": 67, "inclusive": True, "index_base": 0}
    assert unknown["identity_switch_confirmed"] is False
    assert indexed["02-06"]["identity_relationship"] == "same_goalkeeper_reported"
    assert indexed["02-06"]["subject_model_ids"] == [1, 2, 3]
    assert indexed["02-06"]["individual_id_to_range_assignment"] is None
    assert indexed["02-06"]["identity_switch_confirmed"] is False
    assert indexed["12-02"]["review_frame_ranges"][0]["start"] == 145
    assert indexed["12-02"]["review_frame_ranges"][0]["end"] == 170
    assert indexed["12-02"]["whole_clip_pass"] is False
    assert indexed["12-02"]["goalkeeper_coverage_confirmed"] is False
    misses = [row for row in keeper if row["issue_kind"] == "whole_clip_goalkeeper_detection_missing"]
    assert {row["source_video"]["video_id"] for row in misses} == set(importer.FULL_MISSES)
    assert all(row["whole_clip_review_scope"] for row in misses)
    checklist = importer.regression_checklist(keeper + shooter)
    assert len(checklist["items"]) == 20
    assert checklist["model_regression_passed"] is None
    assert checklist["reviewed_items"] == 0
    assert all(row["status"] == "pending" and row["result"] is None for row in checklist["items"])


@pytest.mark.parametrize("start,end", [(-1, 2), (5, 4), (0, 80), (True, 5), (1.0, 5)])
def test_rejects_invalid_frame_ranges(videos, start, end):
    with pytest.raises(ValueError, match="Invalid zero-based"):
        importer.frame_range(start, end, videos["05-04"])


def test_full_clip_miss_must_match_video_extent(source_workbook, videos):
    source, keeper_rows, _ = importer.read_source(source_workbook)
    videos["02-05"]["frame_count"] = 142
    with pytest.raises(ValueError, match="Whole-clip source range differs"):
        importer.keeper_observations(source, keeper_rows, videos)


@pytest.mark.parametrize("change", ["model", "range", "extra_row", "formula", "keeper_observation", "boolean_id"])
def test_rejects_changed_source_contract(source_workbook, change):
    workbook = openpyxl.load_workbook(source_workbook)
    if change == "model":
        workbook[importer.KEEPER_SHEET].cell(2, 1, "YOLOv8-pose")
    elif change == "range":
        workbook[importer.SHOOTER_SHEET].cell(9, 3, 61)
    elif change == "extra_row":
        workbook[importer.SHOOTER_SHEET].cell(19, 1, "05-04")
    elif change == "keeper_observation":
        workbook[importer.KEEPER_SHEET].cell(8, 4, "已确认身份切换")
    elif change == "boolean_id":
        workbook[importer.SHOOTER_SHEET].cell(10, 2, True)
    else:
        workbook[importer.SHOOTER_SHEET].cell(7, 7, "=1+1")
    workbook.save(source_workbook)
    workbook.close()
    with pytest.raises(ValueError):
        importer.read_source(source_workbook)


def test_import_hashes_artifacts_and_refuses_overwrite(source_workbook, videos, tmp_path, monkeypatch):
    video_dir = tmp_path / "videos"
    video_dir.mkdir()
    for clip in videos:
        path = video_dir / f"{clip}.mp4"
        path.write_bytes(clip.encode("ascii"))
        videos[clip].update({"path": str(path), "sha256": importer.sha256(path)})
    monkeypatch.setattr(importer, "video_metadata", lambda path: videos[path.stem])
    output = tmp_path / "imported"
    manifest = importer.import_feedback(source_workbook, video_dir, output)
    assert manifest["counts"]["shooter_issue_ranges"] == 12
    assert manifest["counts"]["coordinate_training_labels"] == 0
    assert manifest["qc_checks"]["accuracy_computed"] is False
    for name, expected in manifest["artifacts"].items():
        assert importer.sha256(output / name) == expected["sha256"]
    saved = json.loads((output / "qc_manifest.json").read_text(encoding="utf-8"))
    assert saved == manifest
    before = {path.name: importer.sha256(path) for path in output.iterdir()}
    with pytest.raises(FileExistsError, match="Refusing to overwrite"):
        importer.import_feedback(source_workbook, video_dir, output)
    assert before == {path.name: importer.sha256(path) for path in output.iterdir()}
